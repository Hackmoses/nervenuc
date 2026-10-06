"""Tests for animal-level group tables and GraphPad/CSV/Excel export preparation."""

from __future__ import annotations

import math
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
import nervenuc.export.writers as export_writers

from nervenuc.aggregation import (
    AnimalAggregationResult,
    SampleIdentity,
    SampleSummary,
    aggregate_by_animal,
)
from nervenuc.export import (
    AnimalGroupAssignment,
    ExportFileExistsError,
    ExportWriteError,
    GroupTableConsistencyError,
    GroupTableResult,
    InvalidGroupAssignmentError,
    prepare_group_tables,
    sanitize_dataframe_for_spreadsheet,
    validate_group_assignments,
    write_group_csv,
    write_group_excel,
    write_dataframe_csv,
)
from nervenuc.summary import (
    BiologicalMetricDefinition,
    BiologicalSummaryResult,
    PhenotypeFilter,
)

EXPECTED_LONG_COLUMNS = [
    "animal_id",
    "group",
    "metric_name",
    "sample_count",
    "numerator_count",
    "denominator_count",
    "percentage",
    "numerator_filter",
    "denominator_filter",
]

EXPECTED_MANIFEST_COLUMNS = [
    "metric_index",
    "metric_name",
    "csv_filename",
    "excel_sheet_name",
    "numerator_filter",
    "denominator_filter",
]


def _create_mock_aggregation_result(
    animal_data: dict[str, list[tuple[int, int]]],
    metric_defs: list[BiologicalMetricDefinition] | None = None,
    sample_count: int = 1,
) -> AnimalAggregationResult:
    """Create a mock AnimalAggregationResult for testing export tables."""
    if metric_defs is None:
        metric_defs = [
            BiologicalMetricDefinition("M1", PhenotypeFilter(), PhenotypeFilter({"BrdU": True})),
            BiologicalMetricDefinition("M2", PhenotypeFilter({"SOX10": True}), PhenotypeFilter({"BrdU": True})),
        ]

    rows = []
    for aid, counts in animal_data.items():
        for (num, den), d in zip(counts, metric_defs):
            pct = float(100.0 * num / den) if den > 0 else float("nan")
            rows.append(
                {
                    "animal_id": aid,
                    "metric_name": d.name,
                    "sample_count": sample_count,
                    "numerator_count": num,
                    "denominator_count": den,
                    "percentage": pct,
                    "numerator_filter": d.effective_numerator.to_string(),
                    "denominator_filter": d.denominator.to_string(),
                }
            )

    df = pd.DataFrame(rows)
    return AnimalAggregationResult(
        animal_metrics=df,
        definitions=tuple(metric_defs),
        samples=(),
    )


def test_single_group_single_animal() -> None:
    """Test minimal valid dataset with 1 animal and 1 group."""
    agg = _create_mock_aggregation_result({"A1": [(5, 10), (2, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]

    res = prepare_group_tables(agg, assigns)

    assert isinstance(res, GroupTableResult)
    assert len(res.long_table) == 2
    assert list(res.long_table.columns) == EXPECTED_LONG_COLUMNS
    assert list(res.long_table["group"]) == ["WT", "WT"]
    assert list(res.long_table["animal_id"]) == ["A1", "A1"]

    # Wide tables
    assert len(res.wide_tables) == 2
    assert "M1" in res.wide_tables and "M2" in res.wide_tables
    m1_df = res.wide_tables["M1"]
    assert list(m1_df.columns) == ["WT"]
    assert m1_df.iloc[0]["WT"] == 50.0

    # Animal ID table
    assert list(res.animal_id_table.columns) == ["WT"]
    assert res.animal_id_table.iloc[0]["WT"] == "A1"


def test_multiple_groups_and_animals_per_group() -> None:
    """Test multiple groups with multiple animals per group."""
    agg = _create_mock_aggregation_result(
        {
            "A1": [(1, 10), (0, 5)],
            "A2": [(2, 10), (1, 5)],
            "A3": [(5, 20), (3, 10)],
            "A4": [(6, 20), (4, 10)],
        }
    )
    assigns = [
        AnimalGroupAssignment("A1", "WT"),
        AnimalGroupAssignment("A2", "WT"),
        AnimalGroupAssignment("A3", "CMT1A"),
        AnimalGroupAssignment("A4", "CMT1A"),
    ]

    res = prepare_group_tables(agg, assigns)

    assert len(res.long_table) == 8  # 4 animals x 2 metrics
    m1_df = res.wide_tables["M1"]
    assert list(m1_df.columns) == ["WT", "CMT1A"]
    assert list(m1_df["WT"]) == [10.0, 20.0]
    assert list(m1_df["CMT1A"]) == [25.0, 30.0]

    assert list(res.animal_id_table["WT"]) == ["A1", "A2"]
    assert list(res.animal_id_table["CMT1A"]) == ["A3", "A4"]


def test_multiple_metrics_manifest_and_tables() -> None:
    """Test multiple metrics produce matching wide tables and manifest rows."""
    d1 = BiologicalMetricDefinition("A_metric", PhenotypeFilter(), PhenotypeFilter({"A": True}))
    d2 = BiologicalMetricDefinition("B_metric", PhenotypeFilter(), PhenotypeFilter({"B": True}))
    d3 = BiologicalMetricDefinition("C_metric", PhenotypeFilter(), PhenotypeFilter({"C": True}))

    agg = _create_mock_aggregation_result(
        {"A1": [(1, 10), (2, 10), (3, 10)]},
        metric_defs=[d1, d2, d3],
    )
    assigns = [AnimalGroupAssignment("A1", "WT")]

    res = prepare_group_tables(agg, assigns)

    assert list(res.wide_tables.keys()) == ["A_metric", "B_metric", "C_metric"]
    assert len(res.manifest) == 3
    assert list(res.manifest.columns) == EXPECTED_MANIFEST_COLUMNS
    assert list(res.manifest["metric_index"]) == [1, 2, 3]
    assert list(res.manifest["metric_name"]) == ["A_metric", "B_metric", "C_metric"]
    assert list(res.manifest["csv_filename"]) == ["metric_001.csv", "metric_002.csv", "metric_003.csv"]
    assert list(res.manifest["excel_sheet_name"]) == ["Metric 001", "Metric 002", "Metric 003"]


def test_long_table_columns_and_counts_preserved() -> None:
    """Test that all long table columns and raw counts are preserved exactly."""
    agg = _create_mock_aggregation_result({"A1": [(15, 30), (5, 10)]}, sample_count=3)
    assigns = [AnimalGroupAssignment("A1", "GroupX")]

    res = prepare_group_tables(agg, assigns)
    row = res.long_table.iloc[0]

    assert row["animal_id"] == "A1"
    assert row["group"] == "GroupX"
    assert row["metric_name"] == "M1"
    assert row["sample_count"] == 3
    assert row["numerator_count"] == 15
    assert row["denominator_count"] == 30
    assert row["percentage"] == 50.0
    assert row["numerator_filter"] == "BrdU+"
    assert row["denominator_filter"] == "ALL"


def test_unequal_group_sizes_padded_with_nan() -> None:
    """Test that groups of unequal sizes are padded with NaN in wide and animal ID tables."""
    agg = _create_mock_aggregation_result(
        {
            "WT1": [(1, 10), (0, 0)],
            "WT2": [(2, 10), (0, 0)],
            "WT3": [(3, 10), (0, 0)],
            "MUT1": [(5, 10), (0, 0)],
        }
    )
    assigns = [
        AnimalGroupAssignment("WT1", "WT"),
        AnimalGroupAssignment("WT2", "WT"),
        AnimalGroupAssignment("WT3", "WT"),
        AnimalGroupAssignment("MUT1", "MUT"),
    ]

    res = prepare_group_tables(agg, assigns)
    m1_df = res.wide_tables["M1"]

    assert len(m1_df) == 3
    assert list(m1_df["WT"]) == [10.0, 20.0, 30.0]
    assert m1_df["MUT"].iloc[0] == 50.0
    assert math.isnan(m1_df["MUT"].iloc[1])
    assert math.isnan(m1_df["MUT"].iloc[2])

    assert list(res.animal_id_table["WT"]) == ["WT1", "WT2", "WT3"]
    assert res.animal_id_table["MUT"].iloc[0] == "MUT1"
    assert pd.isna(res.animal_id_table["MUT"].iloc[1])
    assert pd.isna(res.animal_id_table["MUT"].iloc[2])


def test_undefined_animal_percentage_remains_nan() -> None:
    """Test that 0/0 denominator animal percentage remains NaN without being converted to 0."""
    agg = _create_mock_aggregation_result({"A1": [(0, 0), (2, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]

    res = prepare_group_tables(agg, assigns)
    assert math.isnan(res.long_table.iloc[0]["percentage"])
    assert math.isnan(res.wide_tables["M1"].iloc[0]["WT"])


def test_all_animals_must_be_assigned_completeness() -> None:
    """Test that missing an assignment for any animal in M7 raises InvalidGroupAssignmentError."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)], "A2": [(1, 2), (0, 0)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]

    with pytest.raises(InvalidGroupAssignmentError, match="Missing group assignment"):
        prepare_group_tables(agg, assigns)


def test_unknown_assigned_animal_rejected() -> None:
    """Test that assigning an animal not present in M7 raises InvalidGroupAssignmentError."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]})
    assigns = [AnimalGroupAssignment("A1", "WT"), AnimalGroupAssignment("A99", "WT")]

    with pytest.raises(InvalidGroupAssignmentError, match="Assignment for unknown animal"):
        prepare_group_tables(agg, assigns)


def test_duplicate_identical_and_conflicting_assignments_rejected() -> None:
    """Test that duplicate identical and duplicate conflicting assignments are rejected."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]})

    with pytest.raises(InvalidGroupAssignmentError, match="Duplicate assignment"):
        prepare_group_tables(agg, [AnimalGroupAssignment("A1", "WT"), AnimalGroupAssignment("A1", "WT")])

    with pytest.raises(InvalidGroupAssignmentError, match="Duplicate assignment"):
        prepare_group_tables(agg, [AnimalGroupAssignment("A1", "WT"), AnimalGroupAssignment("A1", "MUT")])


def test_invalid_group_and_animal_ids_rejected() -> None:
    """Test rejection of empty and whitespace-padded group and animal IDs."""
    with pytest.raises(InvalidGroupAssignmentError, match="non-empty string"):
        AnimalGroupAssignment("", "WT")

    with pytest.raises(InvalidGroupAssignmentError, match="non-empty string"):
        AnimalGroupAssignment("A1", "")

    with pytest.raises(InvalidGroupAssignmentError, match="leading or trailing whitespace"):
        AnimalGroupAssignment(" A1", "WT")

    with pytest.raises(InvalidGroupAssignmentError, match="leading or trailing whitespace"):
        AnimalGroupAssignment("A1", "WT ")


def test_exact_group_casing_preserved() -> None:
    """Test that group casing is preserved and not normalized."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)], "A2": [(1, 2), (0, 0)]})
    assigns = [AnimalGroupAssignment("A1", "WT"), AnimalGroupAssignment("A2", "wt")]

    res = prepare_group_tables(agg, assigns)
    assert list(res.wide_tables["M1"].columns) == ["WT", "wt"]


def test_group_ordering_follows_assignment_appearance() -> None:
    """Test that GraphPad wide table columns follow first appearance in assignments."""
    agg = _create_mock_aggregation_result(
        {"A1": [(1, 2), (0, 0)], "A2": [(1, 2), (0, 0)], "A3": [(1, 2), (0, 0)]}
    )
    assigns = [
        AnimalGroupAssignment("A3", "Group_C"),
        AnimalGroupAssignment("A1", "Group_A"),
        AnimalGroupAssignment("A2", "Group_B"),
    ]

    res = prepare_group_tables(agg, assigns)
    assert list(res.wide_tables["M1"].columns) == ["Group_C", "Group_A", "Group_B"]


def test_animal_ordering_within_group_follows_m7() -> None:
    """Test that within each group, animals follow canonical M7 animal order."""
    agg = _create_mock_aggregation_result(
        {"A1": [(1, 2), (0, 0)], "A2": [(1, 2), (0, 0)], "A3": [(1, 2), (0, 0)]}
    )
    # Supplied in reverse assignment order for the same group
    assigns = [
        AnimalGroupAssignment("A3", "WT"),
        AnimalGroupAssignment("A2", "WT"),
        AnimalGroupAssignment("A1", "WT"),
    ]

    res = prepare_group_tables(agg, assigns)
    # Must preserve M7 order: A1, A2, A3
    assert list(res.animal_id_table["WT"]) == ["A1", "A2", "A3"]


def test_m7_sample_count_validation() -> None:
    """Test rejection of non-positive or inconsistent sample_count across metrics."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]}, sample_count=1)

    # Corrupt sample_count to 0
    corrupt_df = agg.animal_metrics.copy()
    corrupt_df.loc[0, "sample_count"] = 0
    corrupt_agg = AnimalAggregationResult(corrupt_df, agg.definitions, agg.samples)

    with pytest.raises(GroupTableConsistencyError, match="sample_count must be >= 1"):
        prepare_group_tables(corrupt_agg, [AnimalGroupAssignment("A1", "WT")])

    # Corrupt sample_count to be inconsistent across metrics
    corrupt_df2 = agg.animal_metrics.copy()
    corrupt_df2.loc[1, "sample_count"] = 2
    corrupt_agg2 = AnimalAggregationResult(corrupt_df2, agg.definitions, agg.samples)

    with pytest.raises(GroupTableConsistencyError, match="Inconsistent sample_count"):
        prepare_group_tables(corrupt_agg2, [AnimalGroupAssignment("A1", "WT")])


def test_m7_duplicate_or_missing_rows_rejected() -> None:
    """Test rejection of missing or duplicate (animal_id, metric_name) rows."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]})

    # Missing row
    short_df = agg.animal_metrics.iloc[:1].copy()
    short_agg = AnimalAggregationResult(short_df, agg.definitions, agg.samples)

    with pytest.raises(GroupTableConsistencyError, match="expected 2"):
        prepare_group_tables(short_agg, [AnimalGroupAssignment("A1", "WT")])


def test_m7_corrupted_counts_or_percentages_rejected() -> None:
    """Test rejection of corrupted percentage in M7 result."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]})
    corrupt_df = agg.animal_metrics.copy()
    corrupt_df.loc[0, "percentage"] = 99.9  # Corrupted (1/2 is 50.0)
    corrupt_agg = AnimalAggregationResult(corrupt_df, agg.definitions, agg.samples)

    with pytest.raises(GroupTableConsistencyError, match="Corrupted percentage"):
        prepare_group_tables(corrupt_agg, [AnimalGroupAssignment("A1", "WT")])


def test_m7_mismatched_filter_provenance_rejected() -> None:
    """Test rejection of altered filter strings in M7 result."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (0, 0)]})
    corrupt_df = agg.animal_metrics.copy()
    corrupt_df.loc[0, "denominator_filter"] = "CORRUPTED"
    corrupt_agg = AnimalAggregationResult(corrupt_df, agg.definitions, agg.samples)

    with pytest.raises(GroupTableConsistencyError, match="Mismatched denominator_filter"):
        prepare_group_tables(corrupt_agg, [AnimalGroupAssignment("A1", "WT")])


def test_write_group_csv_creates_files_and_returns_deterministic_keys(tmp_path: Path) -> None:
    """Test write_group_csv creates all expected CSV files and returns stable dictionary keys."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_dir = tmp_path / "export_csv"
    ret = write_group_csv(res, out_dir, overwrite=False)

    assert set(ret.keys()) == {"long_table", "animal_ids", "manifest", "metric_001", "metric_002"}
    assert (out_dir / "animal_metrics_long.csv").exists()
    assert (out_dir / "animal_ids.csv").exists()
    assert (out_dir / "manifest.csv").exists()
    assert (out_dir / "metric_001.csv").exists()
    assert (out_dir / "metric_002.csv").exists()

    # Verify content
    long_read = pd.read_csv(out_dir / "animal_metrics_long.csv")
    assert len(long_read) == 2
    assert list(long_read.columns) == EXPECTED_LONG_COLUMNS


def test_write_group_csv_preflight_overwrite_protection(tmp_path: Path) -> None:
    """Test that write_group_csv aborts before writing anything if any target file already exists."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_dir = tmp_path / "export_preflight"
    out_dir.mkdir(parents=True)
    # Create one conflicting file
    conflict_file = out_dir / "metric_002.csv"
    conflict_file.write_text("existing content", encoding="utf-8")

    with pytest.raises(ExportFileExistsError, match="overwrite is False"):
        write_group_csv(res, out_dir, overwrite=False)

    # Ensure other files were NOT written
    assert not (out_dir / "animal_metrics_long.csv").exists()
    assert conflict_file.read_text(encoding="utf-8") == "existing content"


def test_write_group_csv_preserves_unrelated_files(tmp_path: Path) -> None:
    """Test that write_group_csv does not delete or touch unrelated files in the output directory."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_dir = tmp_path / "export_unrelated"
    out_dir.mkdir(parents=True)
    unrelated = out_dir / "notes.txt"
    unrelated.write_text("important notes", encoding="utf-8")

    write_group_csv(res, out_dir, overwrite=False)
    assert unrelated.read_text(encoding="utf-8") == "important notes"


def test_write_group_csv_overwrite_true_works(tmp_path: Path) -> None:
    """Test that overwrite=True replaces existing files without error."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_dir = tmp_path / "export_overwrite"
    write_group_csv(res, out_dir, overwrite=False)
    # Overwrite
    write_group_csv(res, out_dir, overwrite=True)
    assert (out_dir / "animal_metrics_long.csv").exists()


def test_write_group_csv_failure_leaves_no_partial_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A staged CSV write failure must not publish a partial coordinated export."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    res = prepare_group_tables(agg, [AnimalGroupAssignment("A1", "WT")])
    out_dir = tmp_path / "failed_export"
    original_to_csv = pd.DataFrame.to_csv
    calls = 0

    def fail_third_write(frame, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 3:
            raise OSError("injected staging failure")
        return original_to_csv(frame, *args, **kwargs)

    monkeypatch.setattr(pd.DataFrame, "to_csv", fail_third_write)
    with pytest.raises(ExportWriteError, match="injected staging failure"):
        write_group_csv(res, out_dir)

    assert not any(out_dir.glob("*.csv"))


def test_atomic_single_csv_failure_preserves_existing_destination(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    target = tmp_path / "existing.csv"
    target.write_text("original\n", encoding="utf-8")

    def fail_write(*args, **kwargs):
        raise OSError("injected single-table failure")

    monkeypatch.setattr(pd.DataFrame, "to_csv", fail_write)
    with pytest.raises(ExportWriteError, match="injected single-table failure"):
        write_dataframe_csv(pd.DataFrame({"value": [1]}), target, overwrite=True)
    assert target.read_text(encoding="utf-8") == "original\n"


def test_export_rejects_forged_group_table(tmp_path: Path) -> None:
    """Export validates mutable result tables against authoritative M7 provenance."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    res = prepare_group_tables(agg, [AnimalGroupAssignment("A1", "WT")])
    res.long_table.loc[0, "percentage"] = 999.0

    with pytest.raises(ExportWriteError, match="integrity validation failed"):
        write_group_csv(res, tmp_path / "forged")


def test_write_group_excel_creates_workbook_with_all_sheets(tmp_path: Path) -> None:
    """Test write_group_excel creates valid .xlsx with all expected sheets."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_file = tmp_path / "results.xlsx"
    ret_path = write_group_excel(res, out_file, overwrite=False)

    assert ret_path == out_file
    assert out_file.exists()

    with pd.ExcelFile(out_file, engine="openpyxl") as excel_file:
        assert excel_file.sheet_names == ["Animal Metrics", "Animal IDs", "Metric 001", "Metric 002", "Manifest"]


def test_csv_and_excel_escape_adversarial_formula_values(tmp_path: Path) -> None:
    """Identifiers and groups must remain inert text in spreadsheet consumers."""
    from openpyxl import load_workbook

    agg = _create_mock_aggregation_result({"=2+2": [(1, 2), (3, 4)]})
    res = prepare_group_tables(agg, [AnimalGroupAssignment("=2+2", "@SUM(A1:A2)")])

    csv_dir = tmp_path / "safe_csv"
    write_group_csv(res, csv_dir)
    csv_text = (csv_dir / "animal_metrics_long.csv").read_text(encoding="utf-8")
    assert "'=2+2" in csv_text
    assert "'@SUM(A1:A2)" in csv_text

    workbook_path = tmp_path / "safe.xlsx"
    write_group_excel(res, workbook_path)
    workbook = load_workbook(workbook_path, data_only=False)
    worksheet = workbook["Animal Metrics"]
    values = [cell for row in worksheet.iter_rows() for cell in row]
    adversarial_cells = [cell for cell in values if isinstance(cell.value, str) and cell.value.startswith("'")]
    assert {cell.value for cell in adversarial_cells} >= {"'=2+2", "'@SUM(A1:A2)"}
    assert all(cell.data_type != "f" for cell in adversarial_cells)


@pytest.mark.parametrize(
    "payload",
    ["=1+1", "+1+1", "-1+1", "@SUM(A1:A2)", "\t=1+1", "\r=1+1", "\n=1+1", "  =1+1"],
)
def test_spreadsheet_sanitizer_neutralizes_all_formula_prefixes(payload: str) -> None:
    safe = sanitize_dataframe_for_spreadsheet(pd.DataFrame({"value": [payload]}))
    assert safe.iloc[0, 0] == "'" + payload


def test_write_group_excel_overwrite_protection(tmp_path: Path) -> None:
    """Test that write_group_excel raises ExportFileExistsError when overwrite is False."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_file = tmp_path / "results_protect.xlsx"
    out_file.write_text("existing", encoding="utf-8")

    with pytest.raises(ExportFileExistsError, match="overwrite is False"):
        write_group_excel(res, out_file, overwrite=False)


def test_write_group_excel_overwrite_true_works(tmp_path: Path) -> None:
    """Test that write_group_excel overwrites when overwrite is True."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_file = tmp_path / "results_overwrite.xlsx"
    write_group_excel(res, out_file, overwrite=False)
    write_group_excel(res, out_file, overwrite=True)
    assert out_file.exists()


def test_write_group_excel_failure_preserves_existing_destination(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Excel is published only after the temporary workbook is complete."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    res = prepare_group_tables(agg, [AnimalGroupAssignment("A1", "WT")])
    out_file = tmp_path / "existing.xlsx"
    out_file.write_bytes(b"original workbook bytes")

    def fail_publish(*args, **kwargs):
        raise OSError("injected workbook failure")

    monkeypatch.setattr(export_writers.os, "replace", fail_publish)
    with pytest.raises(ExportWriteError, match="injected workbook failure"):
        write_group_excel(res, out_file, overwrite=True)

    assert out_file.read_bytes() == b"original workbook bytes"


def test_arbitrary_scientific_metric_names_safe_in_excel(tmp_path: Path) -> None:
    """Test that special characters in metric names (e.g. '/', ':', '*') do not break Excel export."""
    d1 = BiologicalMetricDefinition("SOX10+ / BrdU*", PhenotypeFilter(), PhenotypeFilter({"A": True}))
    agg = _create_mock_aggregation_result({"A1": [(1, 2)]}, metric_defs=[d1])
    assigns = [AnimalGroupAssignment("A1", "WT")]
    res = prepare_group_tables(agg, assigns)

    out_file = tmp_path / "scientific.xlsx"
    write_group_excel(res, out_file, overwrite=False)

    with pd.ExcelFile(out_file, engine="openpyxl") as excel_file:
        assert "Metric 001" in excel_file.sheet_names
        manifest_df = excel_file.parse("Manifest")
        assert manifest_df.iloc[0]["metric_name"] == "SOX10+ / BrdU*"


def test_source_m7_result_immutability() -> None:
    """Test that input M7 aggregation result and DataFrames remain untouched."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    df_copy = agg.animal_metrics.copy(deep=True)
    assigns = [AnimalGroupAssignment("A1", "WT")]

    res = prepare_group_tables(agg, assigns)

    assert agg.animal_metrics.equals(df_copy)
    assert res.long_table is not agg.animal_metrics


def test_deterministic_repeated_table_preparation() -> None:
    """Test that repeated calls to prepare_group_tables yield identical results."""
    agg = _create_mock_aggregation_result({"A1": [(1, 2), (3, 4)]})
    assigns = [AnimalGroupAssignment("A1", "WT")]

    res1 = prepare_group_tables(agg, assigns)
    res2 = prepare_group_tables(agg, assigns)

    assert res1.long_table.equals(res2.long_table)
    assert res1.animal_id_table.equals(res2.animal_id_table)
    assert res1.manifest.equals(res2.manifest)
    for m in res1.wide_tables:
        assert res1.wide_tables[m].equals(res2.wide_tables[m])
